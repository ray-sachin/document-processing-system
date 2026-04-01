"""
DOCX Processor - Handle Microsoft Word documents
"""
from typing import Dict, Any
import io

from app.worker.processors.base import BaseProcessor
from app.models.document import Document


class DocxProcessor(BaseProcessor):
    """Processor for DOCX files."""
    
    def parse(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Parse DOCX file content."""
        try:
            from docx import Document as DocxDocument
            
            docx_file = io.BytesIO(content)
            doc = DocxDocument(docx_file)
            
            # Extract text from paragraphs
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            raw_text = "\n\n".join(paragraphs)
            
            # Extract tables
            tables_data = []
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text for cell in row.cells]
                    table_rows.append(row_cells)
                if table_rows:
                    tables_data.append(table_rows)
            
            # Get core properties
            core_props = {}
            try:
                core_props = {
                    "title": doc.core_properties.title,
                    "author": doc.core_properties.author,
                    "subject": doc.core_properties.subject,
                    "keywords": doc.core_properties.keywords,
                    "created": str(doc.core_properties.created) if doc.core_properties.created else None,
                    "modified": str(doc.core_properties.modified) if doc.core_properties.modified else None,
                }
                core_props = {k: v for k, v in core_props.items() if v}
            except Exception:
                pass
            
            return {
                "raw_text": raw_text,
                "paragraph_count": len(paragraphs),
                "table_count": len(tables_data),
                "tables": tables_data,
                "docx_properties": core_props,
                "has_content": bool(raw_text.strip())
            }
            
        except Exception as e:
            return {
                "raw_text": f"[DOCX parsing failed: {str(e)}]",
                "paragraph_count": 0,
                "table_count": 0,
                "tables": [],
                "docx_properties": {},
                "has_content": False,
                "error": str(e)
            }
    
    def extract(self, parsed_data: Dict[str, Any], document: Document) -> Dict[str, Any]:
        """Extract structured data from DOCX content."""
        raw_text = parsed_data.get("raw_text", "")
        docx_props = parsed_data.get("docx_properties", {})
        
        # Use document properties for title if available
        title = docx_props.get("title") or self._extract_title(raw_text, document.original_filename)
        
        category = self._categorize(raw_text, document.original_filename)
        
        # Use document keywords if available
        doc_keywords = docx_props.get("keywords", "")
        if doc_keywords:
            keywords = [kw.strip() for kw in doc_keywords.split(",")]
        else:
            keywords = self._extract_keywords(raw_text)
        
        summary = self._generate_summary(raw_text)
        
        metadata = {
            "file_type": document.file_type,
            "file_size": document.file_size,
            "paragraph_count": parsed_data.get("paragraph_count", 0),
            "table_count": parsed_data.get("table_count", 0),
            "has_content": parsed_data.get("has_content", False),
            "word_count": len(raw_text.split()) if raw_text else 0,
            "docx_author": docx_props.get("author"),
            "docx_created": docx_props.get("created"),
        }
        
        return {
            "title": title,
            "category": category,
            "summary": summary,
            "keywords": keywords,
            "metadata": metadata,
            "raw_text": raw_text,
            "structured_data": {
                "type": "docx",
                "paragraphs": parsed_data.get("paragraph_count", 0),
                "tables": parsed_data.get("table_count", 0),
                "docx_properties": docx_props
            }
        }
